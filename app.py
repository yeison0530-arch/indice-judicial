import streamlit as st
import pandas as pd
import requests
import base64
import io
import os
import re
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Border, Side, Alignment, Font

# --- CONFIGURACIÓN DE LA PÁGINA ---
st.set_page_config(
    page_title="Generador de Índice Electrónico",
    page_icon="📄",
    layout="wide"
)

# --- FUNCIONES DE LÓGICA ---

def extraer_numero_prefijo(nombre: str):
    """Extrae el número secuencial al inicio de un nombre de archivo (ej. 001Demanda -> 1, '001')."""
    match = re.match(r'^0*(\d+)', nombre.strip())
    if match:
        return int(match.group(1)), match.group(0) # Valor entero y texto capturado
    return None, None

def iniciar_navegador(raw_url: str):
    import time
    from selenium import webdriver
    from selenium.webdriver.edge.options import Options as EdgeOptions
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
    
    options = EdgeOptions()
    options.page_load_strategy = 'eager'
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-software-rasterizer")
    options.add_argument("--mute-audio")
    options.add_experimental_option('excludeSwitches', ['enable-logging'])
    
    try:
        driver = webdriver.Edge(options=options)
        # Mostrar la ventana para que el usuario pueda interactuar
        driver.maximize_window()
        
        driver.set_page_load_timeout(60)
        status = "timeout"
        try:
            driver.get(raw_url)
            # Esperar activamente hasta que el esqueleto de la lista de archivos aparezca en pantalla
            WebDriverWait(driver, 45).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, '.odspSpartanList, [data-automationid="DetailsList"]'))
            )
            time.sleep(2) # Pausa mínima para que el texto de los elementos se termine de renderizar
            status = "ready"
        except Exception:
            pass # Ignoramos timeout si ya cargó el DOM inicial
        return driver, status
    except Exception as e:
        st.error(f"Error iniciando navegador: {e}")
        return None, "error"

def extraer_pantalla_actual(driver, archivos_actuales, nombres_vistos):
    try:
        body_text = driver.execute_script("""
            var lista = document.querySelector('.odspSpartanList') || document.querySelector('[data-automationid="DetailsList"]');
            return lista ? lista.innerText : document.body.innerText;
        """)
        if not body_text: body_text = ""
        lineas = body_text.split('\n')
        
        nuevos_archivos = 0
        for i, linea in enumerate(lineas):
            ln = linea.strip().lower()
            if ln.endswith(" kb") or ln.endswith(" mb") or ln.endswith(" gb") or ln.endswith(" bytes") or ln.endswith(" elementos") or ln.endswith(" elemento"):
                if " " in ln and ln.split(" ")[0].replace(",", "").replace(".", "").isdigit():
                    try:
                        size_str = linea.strip() 
                        mod_date = lineas[i-2].strip()
                        name = lineas[i-3].strip()
                        
                        if name and name not in ["Nombre", "Compartir", "Actividad", "Tamaño del archivo", "Modificado"]:
                            if name not in nombres_vistos:
                                nombres_vistos.add(name)
                                archivos_actuales.append({
                                    "name": name,
                                    "size_str": size_str,
                                    "date_str": mod_date,
                                    "file": "elemento" not in ln
                                })
                                nuevos_archivos += 1
                    except IndexError:
                        pass
        return nuevos_archivos
    except Exception as e:
        st.error(f"Error leyendo la pantalla: {e}")
        return -1

def parse_spanish_date(date_str: str) -> datetime:
    """Traduce formatos humanos (hace 2 horas, 14 abr) a un datetime estricto."""
    date_str = date_str.lower().strip()
    now = datetime.now()
    if not date_str: return now
    
    if "hace" in date_str or "ayer" in date_str:
        if "ayer" in date_str: return now - timedelta(days=1)
        nums = [int(s) for s in date_str.split() if s.isdigit()]
        val = nums[0] if nums else 1
        if "hora" in date_str: return now - timedelta(hours=val)
        elif "minuto" in date_str: return now - timedelta(minutes=val)
        elif "día" in date_str or "dia" in date_str: return now - timedelta(days=val)
        return now
        
    date_str = date_str.replace(" de ", " ")
    if re.match(r"\d{1,2}/\d{1,2}/\d{2,4}", date_str):
        try: return datetime.strptime(date_str, "%d/%m/%Y")
        except: pass
            
    meses_ab = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
    mes_num = 1
    for i, m in enumerate(meses_ab):
        if m in date_str:
            mes_num = i + 1
            break
            
    dia_m = re.search(r"(\d{1,2})", date_str)
    dia = int(dia_m.group(1)) if dia_m else 1
    ano_m = re.search(r"(20\d{2})", date_str)
    ano = int(ano_m.group(1)) if ano_m else now.year
    
    try: return datetime(ano, mes_num, dia)
    except: return now

def format_size(size_bytes: int) -> str:
    """Convierte el tamaño en bytes a un formato legible (KB o MB)"""
    if size_bytes < 1024 * 1024:
        # En Kilobytes
        size_kb = size_bytes / 1024.0
        return f"{size_kb:.2f} KB".replace('.', ',')
    else:
        # En Megabytes
        size_mb = size_bytes / (1024.0 * 1024.0)
        return f"{size_mb:.2f} MB".replace('.', ',')

def process_metadata_to_dataframe(files_data, form_data):
    """Transforma el JSON en un listado apto para el dataframe Oficial con orden estático."""
    # Filtrar carpetas según la meta
    file_items = [item for item in files_data if item.get('file') is True]
    
    # Parsear y asignar claves de ordenamiento
    for item in file_items:
        raw_date = item.get("date_str", "") or item.get("lastModifiedDateTime", "")
        item["dt_obj"] = parse_spanish_date(raw_date)
        num, _ = extraer_numero_prefijo(item.get("name", ""))
        item["sort_num"] = num if num is not None else float('inf')
        
    # Ordenar primero por el número de prefijo (los sin número van al final con inf), y luego por fecha
    file_items.sort(key=lambda x: (x["sort_num"], x["dt_obj"]))
    
    records = []
    pag_global = 1
    
    for idx, item in enumerate(file_items, start=1):
        nombre_doc, ext = os.path.splitext(item.get("name", ""))
        formato = ext.lower().replace(".", "") if ext else "Desconocido"
        
        fecha_str = item["dt_obj"].strftime("%d/%m/%Y")
        num_paginas = 1
        
        record = {
            'Nombre Documento': nombre_doc,
            'Fecha Creación Documento': fecha_str,
            'Fecha Incorporación Expediente': fecha_str,
            'Orden Documento': idx,
            'Número Páginas': num_paginas,
            'Página Inicio': pag_global,
            'Página Fin': pag_global,
            'Formato': formato,
            'Tamaño': item.get("size_str", "Desconocido"),
            'Origen': "Electrónico",
            'Observaciones': form_data.get("observaciones_defecto", "")
        }
        pag_global += 1
        records.append(record)
        
    df = pd.DataFrame(records)
    # Las columnas coinciden con el archivo Excel Maestro (Row 11)
    columnas_oficiales = [
        'Nombre Documento', 'Fecha Creación Documento', 'Fecha Incorporación Expediente',
        'Orden Documento', 'Número Páginas', 'Página Inicio', 'Página Fin',
        'Formato', 'Tamaño', 'Origen', 'Observaciones'
    ]
    if len(df) == 0: return pd.DataFrame(columns=columnas_oficiales)
    return df[columnas_oficiales]

def build_excel_file(df: pd.DataFrame, form_data: dict) -> bytes:
    """Inyecta el DataFrame en la plantilla maestra de IndiceElectrónico.xlsx"""
    template_path = 'IndiceElectrónico.xlsx'
    
    if not os.path.exists(template_path):
        # Fallback normal si el Excel se pierde
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Índice Electrónico')
        return output.getvalue()
        
    wb = openpyxl.load_workbook(template_path)
    sheet = wb.active
    
    # Limpiar posibles excepciones de MergedCells al intentar inyectar en celdas combinadas.
    def write_safe(row, col, value):
        try:
            cell = sheet.cell(row=row, column=col)
            if type(cell).__name__ != 'MergedCell':
                cell.value = value
        except: pass
        
    write_safe(2, 2, form_data.get('ciudad', ''))
    write_safe(3, 2, form_data.get('despacho', ''))
    write_safe(4, 2, form_data.get('serie', ''))
    write_safe(5, 2, form_data.get('radicacion', ''))
    write_safe(6, 2, form_data.get('parte_a', ''))
    write_safe(7, 2, form_data.get('parte_b', ''))
    write_safe(8, 2, form_data.get('terceros', ''))
    write_safe(9, 2, form_data.get('cuaderno', ''))
    
    # Checks manuales (SI / NO) anclados en J3 que en el molde es "SI  X   NO"
    if form_data.get('exped_fisico'):
        write_safe(3, 10, "SI  X    NO")
    else:
        write_safe(3, 10, "SI       NO  X")
        
    # Limpiar filas existentes debajo de las columnas (fila 12+) para inyectar limpio
    if sheet.max_row >= 12:
        sheet.delete_rows(12, sheet.max_row - 11)
        
    # Insertar el DataFrame
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    for idx_row, row in enumerate(df.values, start=12):
        for idx_col, value in enumerate(row, start=1):
            cell = sheet.cell(row=idx_row, column=idx_col)
            cell.value = value
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if idx_col == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center") # Nombre documento align left

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# --- GESTOR DE VALORES PREDETERMINADOS ---
import json
def load_defaults():
    try:
        with open("defaults.json", "r", encoding="utf-8") as f: return json.load(f)
    except:
        return {"ciudad": "Ibagué", "despacho": "Juzgado 08 Civil Municipal", "serie": "PERTENENCIA"}

def save_defaults(data):
    with open("defaults.json", "w", encoding="utf-8") as f: json.dump(data, f)


# --- INTERFAZ DE USUARIO (UI) ---
st.title("Indexador Electrónico")
st.markdown("### REGISTRE LA INFORMACIÓN DEL PROCESO")

defaults = load_defaults()

# Fila 1: Checkboxes
col_cb1, col_cb2, _ = st.columns([1, 1, 2])
with col_cb1: exped_fisico = st.checkbox("Expediente Posee Documentos Físicos")
with col_cb2: act_existente = st.checkbox("Actualizar un Indice Existente")

# Fila 2: Ciudad, Despacho, Predeterminar
c1, c2, c3 = st.columns([1, 2, 1])
with c1: ciudad = st.text_input("Ciudad", value=defaults.get("ciudad", ""))
with c2: despacho = st.text_input("Despacho Judicial", value=defaults.get("despacho", ""))
with c3:
    st.write("") # Espaciador
    guardar_predeterminados = st.checkbox("Predeterminar Valores", value=True)

# Guardado automático de predeterminados
if guardar_predeterminados:
    save_defaults({"ciudad": ciudad, "despacho": despacho})

# Fila 3: Serie y Radicación
c4, c5 = st.columns([2, 2])
with c4: serie = st.text_input("Serie o SubSerie Documental", value=defaults.get("serie", "PERTENENCIA"))
with c5: radicacion = st.text_input("No. Radicación del Proceso")

if guardar_predeterminados: save_defaults({"ciudad": ciudad, "despacho": despacho, "serie": serie})

# Fila 4, 5, 6: Partes
parte_a = st.text_input("Partes Procesales (Parte A) (demandado, procesado, accionado)")
parte_b = st.text_input("Partes Procesales (Parte B) (demandante, denunciante, accionante)")
terceros = st.text_input("Terceros Intervinientes")

# Fila 7: Cuaderno y URL
c6, c7 = st.columns([1, 3])
with c6: cuaderno = st.text_input("Nombre del Cuaderno")
with c7: onedrive_url = st.text_input("🔗 URL de OneDrive", placeholder="Pega el enlace web público aquí...")

# --- ESTADO DE LA SESIÓN PARA EXTRACCIÓN ASISTIDA ---
if 'driver' not in st.session_state:
    st.session_state.driver = None
if 'archivos' not in st.session_state:
    st.session_state.archivos = []
if 'nombres_vistos' not in st.session_state:
    st.session_state.nombres_vistos = set()

st.markdown("---")
st.markdown("### 🤖 EXTRACCIÓN ASISTIDA POR EL USUARIO")

rango_texto = ""
faltantes = []
sin_numerar = 0

if len(st.session_state.archivos) > 0:
    numeros = []
    for a in st.session_state.archivos:
        num, prefix = extraer_numero_prefijo(a["name"])
        if num is not None:
            numeros.append((num, prefix))
        else:
            sin_numerar += 1
            
    if numeros:
        numeros.sort(key=lambda x: x[0])
        min_item = numeros[0]
        max_item = numeros[-1]
        
        esperados = set(range(min_item[0], max_item[0] + 1))
        encontrados = set(x[0] for x in numeros)
        faltantes = sorted(list(esperados - encontrados))
        
        rango_texto = f" (Desde el {min_item[1]} al {max_item[1]})"

if sin_numerar > 0:
    rango_texto += f" y {sin_numerar} archivo(s) sin numerar"

st.info(f"📁 **Archivos extraídos actualmente: {len(st.session_state.archivos)}{rango_texto}**")

if faltantes:
    faltantes_str = ", ".join(map(str, faltantes[:15]))
    if len(faltantes) > 15:
        faltantes_str += "..."
    st.warning(f"⚠️ **Atención:** Parece haber un salto en la secuencia numérica. Faltan los archivos: **{faltantes_str}**. Por favor ubícalos en el navegador y pulsa 'Extraer Más'.")

col_a, col_b, col_c = st.columns(3)

with col_a:
    if st.button("1️⃣ Abrir Navegador (Paso 1)", use_container_width=True):
        if not onedrive_url.strip():
            st.warning("⚠️ Ingresa un enlace de OneDrive.")
        else:
            with st.spinner("⏳ Abriendo Edge y esperando a que SharePoint cargue los archivos..."):
                if st.session_state.driver:
                    try: st.session_state.driver.quit()
                    except: pass
                
                st.session_state.archivos = []
                st.session_state.nombres_vistos = set()
                
                driver, status = iniciar_navegador(onedrive_url.strip())
                if driver:
                    st.session_state.driver = driver
                    if status == "ready":
                        st.success("✅ ¡Carpeta lista y detectada! Ya puedes ir a Edge y pulsar el Paso 2.")
                    else:
                        st.warning("⚠️ El navegador abrió pero tardó en cargar. Verifica en Edge manualmente antes de pulsar el Paso 2.")

with col_b:
    if st.button("2️⃣ Extraer lo que hay en Pantalla", use_container_width=True, type="primary"):
        if not st.session_state.driver:
            st.warning("⚠️ Primero debes iniciar el navegador en el paso 1.")
        else:
            with st.spinner("Leyendo los archivos visibles..."):
                nuevos = extraer_pantalla_actual(st.session_state.driver, st.session_state.archivos, st.session_state.nombres_vistos)
                if nuevos == 0:
                    st.warning("No se encontraron archivos nuevos. Baja un poco más en el navegador.")
                elif nuevos > 0:
                    st.success(f"¡Se sumaron {nuevos} archivos nuevos!")
                    import time
                    time.sleep(1)
                    st.rerun()

with col_c:
    if st.button("3️⃣ Finalizar y Generar Índice", use_container_width=True):
        if len(st.session_state.archivos) == 0:
            st.warning("⚠️ No hay archivos extraídos. Usa los pasos 1 y 2.")
        else:
            if st.session_state.driver:
                try: st.session_state.driver.quit()
                except: pass
                st.session_state.driver = None
                
            form_data = {
                "ciudad": ciudad,
                "despacho": despacho,
                "serie": serie,
                "radicacion": radicacion,
                "parte_a": parte_a,
                "parte_b": parte_b,
                "terceros": terceros,
                "cuaderno": cuaderno,
                "exped_fisico": exped_fisico,
                "act_existente": act_existente
            }
            
            df_indice = process_metadata_to_dataframe(st.session_state.archivos, form_data)
            
            st.subheader("📊 Vista Previa del Índice")
            st.dataframe(df_indice, use_container_width=True, hide_index=True)
            
            excel_data = build_excel_file(df_indice, form_data)
            
            st.download_button(
                label="⬇️ Descargar Índice Oficial (.xlsx)",
                data=excel_data,
                file_name=f"Indice_{radicacion or 'Electronico'}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
