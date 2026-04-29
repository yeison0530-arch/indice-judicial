import openpyxl

wb = openpyxl.load_workbook('IndiceElectrónico.xlsx')
sheet = wb.active

print("Sheet Title:", sheet.title)
print("Max row:", sheet.max_row, "Max col:", sheet.max_column)

for i in range(1, min(15, sheet.max_row + 1)):
    row = []
    for j in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=i, column=j)
        val = cell.value if cell.value is not None else ""
        row.append(str(val))
    print(f"Row {i}:", " | ".join(row))

# Get formatting for the header row (assuming row 6 or something is header)
# Just look for 'Nombre Documento'
header_row = None
for i in range(1, sheet.max_row + 1):
    for j in range(1, sheet.max_column + 1):
        if sheet.cell(row=i, column=j).value == "Nombre Documento":
            header_row = i
            break
    if header_row: break

if header_row:
    print("\nHeaders at row:", header_row)
    for j in range(1, sheet.max_column + 1):
        c = sheet.cell(row=header_row, column=j)
        bg = c.fill.start_color.index if c.fill.start_color else None
        print(f"Col {j} ({c.value}): fill={bg}")
